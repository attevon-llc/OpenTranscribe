"""DOCX renderer (python-docx, MIT): native Word line numbering and page fields.

A DOCX has no fixed pagination: Word or LibreOffice lays it out with whatever fonts the reader
has, so unlike the PDF it cannot promise a stable page:line. What it can do is hand the reader
Word's own line numbering (``w:lnNumType``, restarting per page or continuous) and PAGE/NUMPAGES
fields, and stay editable. The PDF is the citable artifact; the DOCX is the working copy.

Fonts are named, never embedded (no licence surface): ``Noto Sans`` for Latin, ``Noto Naskh
Arabic`` for complex-script (``w:cs``) and ``Noto Sans CJK <region>`` for East Asian text. Word
substitutes whatever is installed when they are missing.

Comments are inline italic paragraphs, not Word comments, so a printout matches the PDF. Every
string passes through :func:`sanitize_for_xml` first: python-docx raises ``ValueError`` on an
XML-1.0-illegal character, and transcripts can contain them.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.section import _Footer
from docx.section import _Header
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from app.services.transcript_document.model import Block
from app.services.transcript_document.model import CommentBlock
from app.services.transcript_document.model import Direction
from app.services.transcript_document.model import DocumentLabels
from app.services.transcript_document.model import DocumentMeta
from app.services.transcript_document.model import DocumentOptions
from app.services.transcript_document.model import PageSpec
from app.services.transcript_document.model import TurnBlock
from app.services.transcript_document.scripts import cjk_font_family
from app.services.transcript_document.scripts import sanitize_for_xml

LATIN_FONT = "Noto Sans"
COMPLEX_SCRIPT_FONT = "Noto Naskh Arabic"
#: Elements the schema places AFTER ``w:bidi`` inside ``w:pPr``.
_PPR_AFTER_BIDI = (
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)
_PLACEHOLDER = re.compile(r"(\{pages\}|\{page\})")
#: Gap between the Word line number and the text, in twips (1/20 pt).
_LINE_NUMBER_DISTANCE_TWIPS = 360


def _set_fonts(rpr: BaseOxmlElement, east_asia: str) -> None:
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:cs"), COMPLEX_SCRIPT_FONT)
    fonts.set(qn("w:eastAsia"), east_asia)


def _style_normal(doc: DocxDocument, spec: PageSpec, east_asia: str) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(spec.font_size)
    rpr = normal.element.get_or_add_rPr()
    _set_fonts(rpr, east_asia)
    # Complex-script text is sized by szCs, not sz; without it Arabic prints at the default.
    size = rpr.find(qn("w:sz"))
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(spec.font_size * 2)))
    size.addnext(size_cs)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)


def _configure_section(doc: DocxDocument, spec: PageSpec, options: DocumentOptions) -> None:
    section = doc.sections[0]
    section.page_width = Pt(spec.width)
    section.page_height = Pt(spec.height)
    section.top_margin = Pt(spec.margin_top)
    section.bottom_margin = Pt(spec.margin_bottom)
    section.left_margin = Pt(spec.gutter_left)
    section.right_margin = Pt(spec.margin_right)
    if options.line_numbering == "none":
        return
    numbering = OxmlElement("w:lnNumType")
    numbering.set(qn("w:countBy"), "1")
    mode = "newPage" if options.line_numbering == "per_page" else "continuous"
    numbering.set(qn("w:restart"), mode)
    numbering.set(qn("w:distance"), str(_LINE_NUMBER_DISTANCE_TWIPS))
    # Schema order: pgSz, pgMar, ..., lnNumType, pgNumType, cols.
    section._sectPr.find(qn("w:pgMar")).addnext(numbering)


def _write(
    paragraph: Paragraph,
    text: str,
    direction: Direction,
    *,
    bold: bool = False,
    italic: bool = False,
    keep_with_next: bool = False,
    space_after: float = 0.0,
) -> Paragraph:
    """Fill ``paragraph`` with one run; a right-to-left one gets ``w:bidi`` and ``w:rtl``."""
    run = paragraph.add_run(sanitize_for_xml(text))
    run.bold = True if bold else None
    run.italic = True if italic else None
    if keep_with_next:
        paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(space_after)
    if direction == "rtl":
        # ``w:bidi`` precedes ``w:spacing`` in the pPr sequence; Word rejects the wrong order.
        paragraph._p.get_or_add_pPr().insert_element_before(OxmlElement("w:bidi"), *_PPR_AFTER_BIDI)
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
    return paragraph


def _add_field(paragraph: Paragraph, instruction: str) -> None:
    """A ``w:fldSimple`` field (PAGE / NUMPAGES) with a placeholder result of ``1``."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f" {instruction} ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _fill_footer(footer: _Footer, label: str) -> None:
    """``page_label`` with ``{page}``/``{pages}`` as live fields; any other brace is literal."""
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for part in _PLACEHOLDER.split(sanitize_for_xml(label)):
        if part == "{page}":
            _add_field(paragraph, "PAGE")
        elif part == "{pages}":
            _add_field(paragraph, "NUMPAGES")
        elif part:
            paragraph.add_run(part)


def _fill_header(header: _Header, meta: DocumentMeta, labels: DocumentLabels) -> None:
    first = header.paragraphs[0]
    _write(first, meta.title, "ltr", bold=True)
    exported = labels.exported_label.replace("{date}", meta.exported_on.isoformat())
    _write(header.add_paragraph(), exported, "ltr")
    if meta.redacted:
        _write(header.add_paragraph(), labels.redacted_label, "ltr", italic=True)


def _add_block(doc: DocxDocument, block: Block) -> None:
    if isinstance(block, CommentBlock):
        _write(doc.add_paragraph(), block.text, block.direction, italic=True, space_after=4)
        return
    assert isinstance(block, TurnBlock)
    if block.header:
        _write(doc.add_paragraph(), block.header, block.direction, bold=True, keep_with_next=True)
    _write(doc.add_paragraph(), block.text, block.direction, space_after=6)


def render_docx(
    blocks: list[Block],
    spec: PageSpec,
    options: DocumentOptions,
    labels: DocumentLabels,
    meta: DocumentMeta,
) -> bytes:
    """The transcript as a ``.docx`` file's bytes."""
    doc = Document()
    _style_normal(doc, spec, cjk_font_family(meta.language))
    _configure_section(doc, spec, options)
    section = doc.sections[0]
    _fill_header(section.header, meta, labels)
    _fill_footer(section.footer, labels.page_label)
    for block in blocks:
        _add_block(doc, block)

    props = doc.core_properties
    props.title = sanitize_for_xml(meta.title)
    # python-docx's template stamps its own name and a 2013 date; this is the user's document.
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.keywords = ""
    props.subject = ""
    props.category = ""
    props.created = meta.exported_at
    props.modified = meta.exported_at

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
